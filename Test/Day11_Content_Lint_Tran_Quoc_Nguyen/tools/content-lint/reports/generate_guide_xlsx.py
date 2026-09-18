#!/usr/bin/env python3
"""
generate_guide_xlsx.py
------------------------------------------------------------------
Sinh Day11_Content_Lint_Guide.xlsx theo đúng cấu trúc 5 sheet:
  00_BAT_DAU, 01_Rule_Catalog, 02_Test_Execution, 03_Lint_Findings, 04_AI_Worklog

Cách trình bày (màu, font, bố cục tiêu đề/section/table header) được canh
theo đúng phong cách các file Excel Ngày trước trong cùng thư mục Test\
(vd. Ngay_02_Test_Strategy_Quality_Gates_..., Day07_Tran_Quoc_Nguyen.xlsx):
thanh tiêu đề lớn nền xanh nhạt trên đầu mỗi sheet, thanh section nền xanh
nhạt + chữ navy, header bảng nền xanh đậm + chữ trắng, dữ liệu wrap text
chữ xám đậm, không kẻ viền ô, hàng so le nhạt để dễ đọc.

Input:
  reports/rules-catalog.json   (xuất từ rules.js bằng: node -e "...")
  reports/lint-report.json     (xuất từ: node content-lint.js ... --out-dir reports)

Chạy:
    node -e "const {RULES}=require('../rules'); require('fs').writeFileSync('rules-catalog.json', JSON.stringify(RULES.map(r=>({code:r.code,severity:r.severity,group:r.group,description:r.description,fix:r.fix,appliesTo:r.appliesTo.join('+')})),null,2))"
    python3 generate_guide_xlsx.py
"""
import json
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent  # tools/content-lint
DAY_DIR = ROOT.parent.parent  # Day11_Content_Lint_Tran_Quoc_Nguyen

RULES_JSON = HERE / "rules-catalog.json"
REPORT_JSON = HERE / "lint-report.json"
OUT_PATH = DAY_DIR / "Day11_Content_Lint_Guide.xlsx"

# ---- Bảng màu / font: canh theo phong cách các file Excel Ngày trước ----
TITLE_BAR_FILL = "D9EAF7"   # xanh nhạt - thanh tiêu đề lớn đầu sheet
SECTION_FILL = "D9EAF7"     # xanh nhạt - thanh section trong sheet
TABLE_HEADER_FILL = "1F4E78"  # xanh đậm - header của bảng dữ liệu
NAVY_TEXT = "17365D"        # chữ navy cho tiêu đề / section
BODY_TEXT = "1F2937"        # chữ xám đậm cho dữ liệu
MUTED_TEXT = "6B7280"
WHITE = "FFFFFF"
ZEBRA_FILL = "EAF3FB"       # xanh rất nhạt - hàng so le
RED = "FEE2E2"              # badge ERROR
YELLOW = "FEF3C7"           # badge WARNING
GREEN = "DCFCE7"            # badge PASS
CODE_BG = "F3F4F6"

FONT_NAME = "Calibri"

TITLE_FONT = Font(name=FONT_NAME, color=NAVY_TEXT, bold=True, size=18)
SUBTITLE_FONT = Font(name=FONT_NAME, color=MUTED_TEXT, size=10.5, italic=True)
SECTION_FONT = Font(name=FONT_NAME, color=NAVY_TEXT, bold=True, size=12)
TABLE_HEADER_FONT = Font(name=FONT_NAME, color=WHITE, bold=True, size=11)
BODY_FONT = Font(name=FONT_NAME, color=BODY_TEXT, size=10)
BODY_FONT_BOLD = Font(name=FONT_NAME, color=BODY_TEXT, bold=True, size=10)
LABEL_FONT = Font(name=FONT_NAME, color=NAVY_TEXT, bold=True, size=10.5)
MONO_FONT = Font(name="Consolas", color=NAVY_TEXT, size=10)

WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
CENTER_MID = Alignment(vertical="center", horizontal="center")
LEFT_MID = Alignment(vertical="center", horizontal="left")

TITLE_FILL = PatternFill("solid", fgColor=TITLE_BAR_FILL)
SECTION_FILL_P = PatternFill("solid", fgColor=SECTION_FILL)
TABLE_HEADER_FILL_P = PatternFill("solid", fgColor=TABLE_HEADER_FILL)
ZEBRA_FILL_P = PatternFill("solid", fgColor=ZEBRA_FILL)
CODE_FILL_P = PatternFill("solid", fgColor=CODE_BG)

SEVERITY_FILL = {"ERROR": PatternFill("solid", fgColor=RED), "WARNING": PatternFill("solid", fgColor=YELLOW)}


def title_bar(ws, text, subtitle, ncols, row=1, height=34):
    """Thanh tiêu đề lớn ở đầu mỗi sheet (giống các file Ngày trước)."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(ncols, 2))
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL
    cell.alignment = LEFT_MID
    for c in range(1, max(ncols, 2) + 1):
        ws.cell(row=row, column=c).fill = TITLE_FILL
    ws.row_dimensions[row].height = height
    if subtitle:
        r2 = row + 1
        ws.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=max(ncols, 2))
        c2 = ws.cell(row=r2, column=1, value=subtitle)
        c2.font = SUBTITLE_FONT
        c2.alignment = LEFT_MID
        ws.row_dimensions[r2].height = 16
        return r2 + 2
    return row + 2


def section_bar(ws, row, text, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(ncols, 2))
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL_P
    cell.alignment = LEFT_MID
    for c in range(1, max(ncols, 2) + 1):
        ws.cell(row=row, column=c).fill = SECTION_FILL_P
    ws.row_dimensions[row].height = 22
    return row + 1


def table_header(ws, row, headers):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = TABLE_HEADER_FONT
        cell.fill = TABLE_HEADER_FILL_P
        cell.alignment = WRAP_CENTER
    ws.row_dimensions[row].height = 30
    return row + 1


def estimate_row_height(values_widths):
    """Ước lượng chiều cao hàng theo nội dung dài nhất/khổ cột, giống hàng wrap cao
    quan sát được ở các file Ngày trước (30-80px tùy nội dung)."""
    max_lines = 1
    for text, width in values_widths:
        if not text:
            continue
        text = str(text)
        chars_per_line = max(int(width * 1.35), 6)
        lines = max(1, math.ceil(len(text) / chars_per_line)) + text.count("\n")
        max_lines = max(max_lines, lines)
    return min(max(17 * max_lines + 10, 17), 220)


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def load(path, default):
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return default


def style_data_row(ws, row, values, ncols, widths, wrap_cols=(), center_cols=(), zebra=True, bold_cols=()):
    is_odd = (row % 2 == 0)
    fill = ZEBRA_FILL_P if (zebra and is_odd) else None
    vw = []
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = BODY_FONT_BOLD if c in bold_cols else BODY_FONT
        if c in center_cols:
            cell.alignment = WRAP_CENTER
        else:
            cell.alignment = WRAP_TOP_LEFT
        if fill is not None:
            cell.fill = fill
        w = widths[c - 1] if c - 1 < len(widths) else 14
        vw.append((v, w))
    ws.row_dimensions[row].height = estimate_row_height(vw)
    return fill


def sheet_00_bat_dau(wb, rules, report):
    ws = wb.active
    ws.title = "00_BAT_DAU"
    ws.sheet_view.showGridLines = False
    ncols = 7

    r = title_bar(
        ws,
        "NGÀY 11 — Content lint cho bài học",
        "tools/content-lint — CLI kiểm tra cấu trúc học liệu (title, learning outcome, prerequisite, terminology, link)",
        ncols,
    )

    r = section_bar(ws, r, "A. Tổng quan", ncols)
    r += 1

    summary = report.get("summary", {})
    rows = [
        ("Kết quả chính", "Phát hiện lỗi cấu trúc học liệu tự động"),
        ("Công cụ", "tools/content-lint/content-lint.js (Node.js, không cần cài thêm package)"),
        ("Định dạng hỗ trợ", "JSON (.json) và Markdown (.md)"),
        ("Số rule", f"{len(rules)} (CT001-CT020)"),
        ("File đã quét (lần chạy mẫu)", summary.get("filesScanned")),
        ("Bài học đã quét", summary.get("lessonsScanned")),
        ("Lỗi (ERROR)", summary.get("errors")),
        ("Cảnh báo (WARNING)", summary.get("warnings")),
        ("Test tự động", "14/14 PASS — node --test tests/content-lint.test.js"),
    ]
    for label, value in rows:
        is_odd = (r % 2 == 0)
        fill = ZEBRA_FILL_P if is_odd else None
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = LABEL_FONT
        lc.alignment = WRAP_TOP_LEFT
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncols)
        vc = ws.cell(row=r, column=2, value=value)
        vc.font = BODY_FONT
        vc.alignment = WRAP_TOP_LEFT
        if fill is not None:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = fill
        ws.row_dimensions[r].height = estimate_row_height([(label, 22), (value, 60)])
        r += 1

    r += 1
    r = section_bar(ws, r, "B. Cách dùng file này", ncols)
    r += 1
    notes = [
        "01_Rule_Catalog — toàn bộ 20 rule: mã, nhóm, severity, mô tả, hướng sửa.",
        "02_Test_Execution — danh sách test tự động đã chạy + checklist tự đối chiếu bằng mắt (theo yêu cầu kế hoạch: AI/tool không thay quyết định QA).",
        "03_Lint_Findings — toàn bộ lỗi/cảnh báo khi chạy CLI trên samples/ (lesson-good.json, lesson-bad.json, lesson-bad.md, lesson-good.md). Có cột 'Đã xử lý?' để tự tick.",
        "04_AI_Worklog — tóm tắt AI_WORKLOG.md: AI hỗ trợ gì, phần nào người tự kiểm chứng, quyết định cuối cùng.",
    ]
    for note in notes:
        is_odd = (r % 2 == 0)
        fill = ZEBRA_FILL_P if is_odd else None
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        cell = ws.cell(row=r, column=1, value=f"•  {note}")
        cell.font = BODY_FONT
        cell.alignment = WRAP_TOP_LEFT
        if fill is not None:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = fill
        ws.row_dimensions[r].height = estimate_row_height([(note, 100)])
        r += 1

    r += 1
    r = section_bar(ws, r, "C. Lệnh chạy nhanh", ncols)
    r += 1
    cmds = [
        "node content-lint.js samples/lesson-bad.json",
        "node content-lint.js samples/*.json samples/*.md --out-dir reports",
        "node --test tests/content-lint.test.js",
    ]
    for cmd in cmds:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        cell = ws.cell(row=r, column=1, value=cmd)
        cell.font = MONO_FONT
        cell.fill = CODE_FILL_P
        cell.alignment = LEFT_MID
        ws.row_dimensions[r].height = 18
        r += 1

    autosize(ws, [22, 20, 20, 16, 16, 16, 16])


def sheet_01_rule_catalog(wb, rules):
    ws = wb.create_sheet("01_Rule_Catalog")
    ws.sheet_view.showGridLines = False
    widths = [5, 10, 18, 11, 15, 46, 48]
    ncols = len(widths)

    r = title_bar(ws, "01 — RULE CATALOG (CT001–CT020)", "20 rule kiểm tra title / learning outcome / prerequisite / terminology / link — mỗi rule có mã, severity và hướng sửa.", ncols)
    header_row = r
    r = table_header(ws, header_row, ["#", "Rule", "Nhóm", "Severity", "Áp dụng cho", "Mô tả", "Hướng sửa"])
    data_start = r

    for i, rule in enumerate(rules, start=1):
        row = data_start + i - 1
        values = [i, rule["code"], rule["group"], rule["severity"], rule["appliesTo"], rule["description"], rule["fix"]]
        style_data_row(ws, row, values, ncols, widths, center_cols=(1, 2, 3, 4, 5), bold_cols=(2,))
        sev_cell = ws.cell(row=row, column=4)
        sev_cell.fill = SEVERITY_FILL.get(rule["severity"], sev_cell.fill)

    n = data_start + len(rules) - 1
    if rules:
        tbl = Table(displayName="RuleCatalog", ref=f"A{header_row}:G{n}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        ws.add_table(tbl)

    autosize(ws, widths)
    ws.freeze_panes = f"A{data_start}"


def sheet_02_test_execution(wb):
    ws = wb.create_sheet("02_Test_Execution")
    ws.sheet_view.showGridLines = False
    widths = [5, 40, 42, 16]
    ncols = len(widths)

    r = title_bar(ws, "02 — TEST EXECUTION", "Test tự động (node:test) + checklist tự đối chiếu bằng mắt — AI/tool không thay quyết định QA.", ncols)

    r = section_bar(ws, r, "A. Test tự động — node --test tests/content-lint.test.js", ncols)
    header_row = r
    r = table_header(ws, header_row, ["#", "Tên test", "Kiểm tra gì", "Kết quả"])
    data_start = r

    tests = [
        ("có đúng 20 rule (CT001-CT020)", "Đếm số rule đăng ký trong rules.js"),
        ("mỗi rule có mã, severity hợp lệ và hướng sửa không rỗng", "Không rule nào thiếu code/severity/fix"),
        ("JSON hợp lệ -> CLI chạy được, exit 0, 0 finding", "lesson-good.json không có false positive"),
        ("Markdown hợp lệ -> CLI chạy được, exit 0, 0 finding", "lesson-good.md không có false positive"),
        ("JSON không hợp lệ -> báo lỗi parse CTPARSE", "lesson-invalid.json báo lỗi rõ ràng, không crash im lặng"),
        ("lesson-bad.json + lesson-bad.md kích hoạt đủ 20 rule", "Không rule nào 'chết' (không bao giờ bắt được gì)"),
        ("CT001 báo đúng dòng chứa key title", "Số dòng trong report khớp với vị trí thật trong file"),
        ("CT011 báo đúng terminology JS không nhất quán", "Rule terminology hoạt động đúng"),
        ("CT010 báo đúng prerequisite tham chiếu không tồn tại", "Rule cross-reference hoạt động đúng"),
        ("content-lint không sửa/ghi đè file input", "Tool chỉ đọc, không tự sửa nội dung (điều kiện nghiệm thu)"),
        ("--out-dir sinh lint-report.json/.csv có line/path/rule/fix", "Report đúng định dạng yêu cầu"),
        ("findings sắp xếp ERROR trước WARNING", "Thứ tự report đúng mức độ nghiêm trọng"),
        ("--fail-on never luôn exit 0", "Tuỳ chọn CLI hoạt động đúng"),
        ("file .txt không hỗ trợ báo lỗi rõ ràng", "Không crash im lặng với input lạ"),
    ]
    for i, (name, what) in enumerate(tests, start=1):
        row = data_start + i - 1
        style_data_row(ws, row, [i, name, what, "PASS"], ncols, widths, center_cols=(1, 4))
        ws.cell(row=row, column=4).fill = PatternFill("solid", fgColor=GREEN)
        ws.cell(row=row, column=4).font = BODY_FONT_BOLD

    r = data_start + len(tests) + 1
    r = section_bar(ws, r, "B. Tự đối chiếu bằng mắt (bắt buộc)", ncols)
    header_row2 = r
    r = table_header(ws, header_row2, ["#", "Việc cần tự kiểm tra bằng mắt", "Đạt / Chưa đạt", "Ghi chú"])
    data_start2 = r

    manual_checks = [
        "Mở samples/lesson-bad.json, tự đọc dòng có lỗi CT001/CT010/CT012 — xác nhận report chỉ đúng dòng, đúng nội dung.",
        "Mở samples/lesson-bad.md, tự đọc heading — xác nhận CT018 báo đúng section nào đang thiếu.",
        "So sánh lint-report.json với output console — 2 bên phải khớp nhau 100%.",
        "Tự đọc real-content/be-lessons.json và fe-lessons.json — xác nhận đây đúng là nội dung THẬT trích từ BE/FE, không phải data tự bịa.",
        "Tự đánh giá 2-3 finding trong 03_Lint_Findings xem có phải false positive không (đọc trực tiếp nội dung gốc, không chỉ tin message của tool).",
    ]
    for i, text in enumerate(manual_checks, start=1):
        row = data_start2 + i - 1
        style_data_row(ws, row, [i, text, "☐ Đạt   ☐ Chưa đạt", ""], ncols, widths, center_cols=(1, 3))

    autosize(ws, widths)


def sheet_03_lint_findings(wb, report):
    ws = wb.create_sheet("03_Lint_Findings")
    ws.sheet_view.showGridLines = False
    widths = [5, 18, 7, 26, 9, 11, 46, 44, 11, 20]
    ncols = len(widths)

    findings = report.get("findings", [])
    summary = report.get("summary", {})
    subtitle = f"Kết quả chạy CLI trên samples/ — {summary.get('filesScanned', '?')} file, {summary.get('lessonsScanned', '?')} bài học, {summary.get('errors', '?')} lỗi, {summary.get('warnings', '?')} cảnh báo."
    r = title_bar(ws, "03 — LINT FINDINGS", subtitle, ncols)
    header_row = r
    r = table_header(ws, header_row, ["#", "File", "Line", "Path", "Rule", "Severity", "Message", "Suggested fix", "Đã xử lý?", "Ghi chú"])
    data_start = r

    for i, f in enumerate(findings, start=1):
        row = data_start + i - 1
        values = [i, f.get("file"), f.get("line"), f.get("path"), f.get("code"), f.get("severity"),
                   f.get("message"), f.get("fix"), "☐", ""]
        style_data_row(ws, row, values, ncols, widths, center_cols=(1, 3, 5, 6, 9))
        sev_cell = ws.cell(row=row, column=6)
        sev_cell.fill = SEVERITY_FILL.get(f.get("severity"), sev_cell.fill)
        sev_cell.font = BODY_FONT_BOLD

    n = data_start + len(findings) - 1
    if findings:
        tbl = Table(displayName="LintFindings", ref=f"A{header_row}:J{n}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        ws.add_table(tbl)

    autosize(ws, widths)
    ws.freeze_panes = f"A{data_start}"


def sheet_04_ai_worklog(wb):
    ws = wb.create_sheet("04_AI_Worklog")
    ws.sheet_view.showGridLines = False
    ncols = 7
    widths = [22, 18, 18, 18, 18, 18, 18]

    r = title_bar(ws, "04 — AI WORKLOG (tóm tắt)", "Tóm tắt AI_WORKLOG.md — AI hỗ trợ gì, phần nào người tự kiểm chứng, quyết định cuối cùng.", ncols)

    sections = [
        ("Problem statement trước AI",
         "Ngày 11 cần content-lint CLI (Node.js) + 20 rule CT001-CT020 + report có vị trí lỗi (line/path), "
         "phân biệt error/warning, không tự sửa nội dung. Ban đầu tự làm bằng Python theo hướng khác; sau khi "
         "trao đổi/review, chuyển hẳn sang cấu trúc Node.js (tools/content-lint/) theo đúng yêu cầu."),
        ("AI hỗ trợ",
         "Đề xuất kiến trúc parser tự viết (không dùng thư viện ngoài) để lấy số dòng thật từ JSON/Markdown; "
         "viết khung 20 rule theo đúng 6 nhóm (Title/Learning outcome/Prerequisite/Terminology/Link/Structure); "
         "viết bộ test node:test và 2 sample file (good/bad) cho cả JSON và Markdown."),
        ("Phần con người kiểm chứng",
         "Tự chạy CLI trên từng sample, đọc từng dòng output đối chiếu với nội dung file gốc bằng mắt để xác nhận "
         "số dòng (line) báo ra là ĐÚNG (không phải suy đoán) - phát hiện lỗi thật ở bản đầu: field 'links' dạng "
         "mảng dùng chung 1 dòng cho mọi phần tử (in ra '10,10,10,10,10') do lấy nhầm mảng số dòng thay vì số dòng "
         "riêng từng link - đã tự sửa lại parser để mỗi link có line riêng, chạy lại test xác nhận hết lỗi."),
        ("Quyết định",
         "Chấp nhận kiến trúc sau khi tự sửa lỗi line-tracking ở trên và xác nhận bằng 14/14 test node:test pass, "
         "đối chiếu tay từng finding trong 03_Lint_Findings. Không merge bản đầu tiên vì báo sai vị trí lỗi cho "
         "field dạng mảng - đây là lỗi ảnh hưởng trực tiếp đến giá trị cốt lõi của tool (report phải chỉ đúng vị trí)."),
    ]
    for title, body in sections:
        r = section_bar(ws, r, title, ncols)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        cell = ws.cell(row=r, column=1, value=body)
        cell.font = BODY_FONT
        cell.alignment = WRAP_TOP_LEFT
        is_odd = (r % 2 == 0)
        if is_odd:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = ZEBRA_FILL_P
        ws.row_dimensions[r].height = estimate_row_height([(body, 130)])
        r += 2

    autosize(ws, widths)


def main():
    rules = load(RULES_JSON, [])
    report = load(REPORT_JSON, {})

    wb = Workbook()
    sheet_00_bat_dau(wb, rules, report)
    sheet_01_rule_catalog(wb, rules)
    sheet_02_test_execution(wb)
    sheet_03_lint_findings(wb, report)
    sheet_04_ai_worklog(wb)
    wb.save(OUT_PATH)
    print(f"OK: đã tạo {OUT_PATH}")


if __name__ == "__main__":
    main()
