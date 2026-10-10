#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kiểm chứng ĐỘC LẬP Defect Dashboard Ngày 23 (không import defect_data.py).

Đọc file Excel đã sinh, đếm lại từ sheet Defect_Log, mở artefact gốc để đối chiếu một số con số,
và kiểm 3 điều kiện nghiệm thu: phân biệt symptom/root cause, action có owner/deadline, không dùng số liệu giả.
Chạy: python tools/verify_defect_dashboard.py [file_da_tinh_lai.xlsx]   (exit code 0 = tất cả PASS)
"""
import collections, csv, difflib, glob, json, os, re, sys
from urllib.parse import unquote
import openpyxl

DAY = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); TEST = os.path.dirname(DAY)
XLSX = os.environ.get("DAY23_XLSX") or os.path.join(DAY, "Day23_Defect_Dashboard_Tran_Quoc_Nguyen.xlsx")
res = []


def check(name, ok, detail=""):
    res.append(bool(ok)); print("%-4s %s%s" % ("PASS" if ok else "FAIL", name, (" -> " + str(detail)) if detail not in ("", None) else ""))


def table(ws, hrow):
    head = [c.value for c in ws[hrow]]; out = []
    for row in ws.iter_rows(min_row=hrow + 1):
        if row[0].value is None:
            break
        out.append({h: c for h, c in zip(head, row) if h})
    return out


wb = openpyxl.load_workbook(XLSX)
log = table(wb["Defect_Log"], 1)
V = lambda r, k: r[k].value   # noqa: E731
tax = [c.value for row in wb["Taxonomy"].iter_rows(min_col=1, max_col=1) for c in row if c.value]

print("== A. Cấu trúc và taxonomy")
check("Đủ 8 sheet theo yêu cầu", wb.sheetnames == ["README", "Defect_Log", "Taxonomy", "Dashboard", "Pareto", "RCA_Top5", "Prevention_Actions", "AI_WORKLOG"], wb.sheetnames)
ids = [V(r, "Defect ID") for r in log]
check("Defect ID duy nhất (%d defect)" % len(ids), len(ids) == len(set(ids)))
need = ["Source Day", "Component", "Summary", "Symptom (hiện tượng)", "Category", "Severity", "Status", "Root Cause (nguyên nhân gốc)",
        "Root Cause Category", "Detection Method", "Owner Team (đề xuất)", "Detected Date", "Evidence", "Defect Kind"]
empty = [(V(r, "Defect ID"), k) for r in log for k in need if V(r, k) in (None, "")]
check("Không có ô bắt buộc nào trống", not empty, empty[:5])
for col in ("Category", "Severity", "Status", "Root Cause Category", "Defect Kind", "Detection Phase"):
    bad = sorted({V(r, col) for r in log} - set(tax))
    check("Mọi giá trị cột %s nằm trong Taxonomy" % col, not bad, bad)

print("== B. Nghiệm thu 1 - phân biệt symptom và root cause")
same = []
for r in log:
    s, c = str(V(r, "Symptom (hiện tượng)")).lower(), str(V(r, "Root Cause (nguyên nhân gốc)")).lower()
    if difflib.SequenceMatcher(None, s, c).ratio() > 0.6 or c in s or s in c:
        same.append(V(r, "Defect ID"))
check("Root cause không lặp lại symptom", not same, same)
short = [V(r, "Defect ID") for r in log if len(str(V(r, "Root Cause (nguyên nhân gốc)"))) < 30]
check("Root cause đủ cụ thể (>= 30 ký tự)", not short, short)
unsure = [r for r in log if re.search(r"chưa xác định|không xác định|nghi ", str(V(r, "Root Cause (nguyên nhân gốc)")).lower())]
bad = [V(r, "Defect ID") for r in unsure if V(r, "Status") != "Needs verification" and V(r, "Defect Kind") != "External"]
check("Root cause còn là giả thuyết thì Status = Needs verification (%d dòng)" % len(unsure), not bad, bad)

print("== C. Nghiệm thu 3 - không dùng số liệu giả")
missing = []
for r in log:
    link = r["Mở"].hyperlink; target = link.target if link else None
    if not target:
        missing.append(V(r, "Defect ID")); continue
    if target.startswith("http"):
        continue
    if not os.path.exists(os.path.normpath(os.path.join(DAY, unquote(target)))):
        missing.append(V(r, "Defect ID"))
check("Mọi defect có evidence mở được (file tồn tại hoặc URL)", not missing, missing)
bad = [V(r, "Defect ID") for r in log if not re.fullmatch(r"\d{4}-\d\d-\d\d|Not recorded", str(V(r, "Detected Date")))]
check("Detected Date là ngày thật hoặc 'Not recorded'", not bad, bad)
bad = [V(r, "Defect ID") for r in log if V(r, "Deadline") not in (None, "", "Chưa xác nhận")]
check("Không có deadline tự đặt cho defect", not bad, bad)
bad = [V(r, "Defect ID") for r in log if V(r, "Fixed Date") not in (None, "", "Not recorded", "2026-10-05")]
bad += [V(r, "Defect ID") for r in log if V(r, "Last Verified") not in (None, "", "2026-10-05", "2026-10-07")]
check("Fixed Date chỉ ghi khi có bằng chứng, còn lại 'Not recorded'", not bad, bad)
bad = [V(r, "Defect ID") for r in log if V(r, "Fixed Date") and V(r, "Status") not in ("Fixed", "Closed", "Retest")]
check("Chỉ defect Fixed/Closed/Retest mới có Fixed Date", not bad, bad)

# đối chiếu vài con số với artefact gốc, đếm bằng cách khác
x = open(glob.glob(os.path.join(TEST, "Day07_Authentication_Tests*/reports/junit.xml"))[0], encoding="utf-8").read()
check("Ngày 7: junit có đúng 1 test fail (DEF-001)", len(re.findall(r"<(failure|error)\b", x)) == 1)
x = open(os.path.join(TEST, "Day08_UI_Smoke/reports/regressions/junit.xml"), encoding="utf-8").read()
check("Ngày 8: regression junit có 2 test fail (DEF-002, DEF-003)", len(re.findall(r"<(failure|error)\b", x)) == 2)
j = json.load(open(os.path.join(TEST, "Day11_Content_Lint_Tran_Quoc_Nguyen/tools/content-lint/real-content/lint-report.json"), encoding="utf-8"))
f = j.get("findings", j) if isinstance(j, dict) else j
check("Ngày 11: 2 ERROR CT004 trên dữ liệu thật (DEF-012)", sum(1 for i in f if i.get("severity") == "ERROR") == 2)
ws = openpyxl.load_workbook(glob.glob(os.path.join(TEST, "day_16_toi_day21_test/day19/*.xlsx"))[0], data_only=True).worksheets[0]
rows = [r for r in ws.iter_rows(values_only=True) if r and any(isinstance(c, str) and re.fullmatch(r"RT-\d+", c) for c in r)]
fails = [r for r in rows if "FAIL" in r]
check("Ngày 19: 13 case FAIL = 10 CHILD_SAFETY (DEF-033) + 3 case riêng (DEF-034..036)",
      len(fails) == 13 and sum(1 for r in fails if "CHILD_SAFETY" in r) == 10, (len(fails),))
ws = openpyxl.load_workbook(glob.glob(os.path.join(TEST, "day_16_toi_day21_test/day20/*.xlsx"))[0], data_only=True)["Calibration"]
head = next(r for r in ws.iter_rows(values_only=True) if r and "Case ID" in r); col = head.index("Agreement v1")
vals = [r[col] for r in ws.iter_rows(values_only=True) if r[col] in ("AGREE", "DISAGREE")]
check("Ngày 20: Judge v1 lệch 11/20 (DEF-037)", vals.count("DISAGREE") == 11 and len(vals) == 20)
ev = "".join(open(f, encoding="utf-8").read() for f in sorted(glob.glob(os.path.join(DAY, "evidence", "live_test_*.md"))))
n_new = len(re.findall(r"^\| N-\d\d \|", ev, re.M)); d23 = [r for r in log if V(r, "Source Day") == 23]
d23web = [r for r in d23 if V(r, "Detection Phase") == "Manual test"]
check("Ngày 23: 13 quan sát N-01..N-13 trong 2 file evidence, gộp thành %d defect" % len(d23web), n_new == 13 and len(d23web) == 11, (n_new, len(d23web)))
check("File evidence không chứa access token", "eyJ" not in ev)

print("== D. Mọi số trên Dashboard/Pareto truy được về Defect_Log")
dash = wb["Dashboard"]
nums = [c for row in dash.iter_rows() for c in row if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
keys = [c for c in nums if c.column != 1]
check("Dashboard không có số nhập tay ở cột giá trị", not keys, [c.coordinate for c in keys][:5])
forms = [c.value for row in dash.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
check("Mọi công thức Dashboard đọc từ Defect_Log (%d công thức)" % len(forms), forms and all("Defect_Log!" in f for f in forms))
par = []
for _r in table(wb["Pareto"], 4):   # dừng ở dòng Tổng, bỏ dòng ghi chú bên dưới
    if V(_r, "Root Cause Category") == "Tổng":
        break
    par.append(_r)
check("Pareto: số lỗi là COUNTIF trên Defect_Log", all(str(V(r, "Số lỗi")).startswith("=COUNTIF(Defect_Log!") for r in par))
cnt = collections.Counter(V(r, "Root Cause Category") for r in log)
order = [V(r, "Root Cause Category") for r in par]
check("Pareto: đủ mọi root cause và sắp xếp giảm dần", sorted(order) == sorted(cnt) and [cnt[o] for o in order] == sorted(cnt.values(), reverse=True), [cnt[o] for o in order])

expected = {"Tổng số defect": len(log), "Fixture bug": sum(1 for r in log if V(r, "Fixture Bug") == "Y"), "Regression": sum(1 for r in log if V(r, "Regression") == "Y")}
st = collections.Counter(V(r, "Status") for r in log)
expected["Đang mở (Open + In Progress)"] = st["Open"] + st["In Progress"]; expected["Đã sửa/đã đóng (Fixed + Closed)"] = st["Fixed"] + st["Closed"]
print("     Giá trị đếm lại từ Defect_Log:", expected, "| status:", dict(st), "| pareto:", [cnt[o] for o in order])
if len(sys.argv) > 1:   # file đã được Excel/LibreOffice tính lại công thức
    w2 = openpyxl.load_workbook(sys.argv[1], data_only=True)["Dashboard"]
    got = {str(r[0]): r[1] for r in w2.iter_rows(min_row=5, max_row=16, max_col=2, values_only=True)}
    check("Giá trị công thức sau khi tính lại khớp số đếm độc lập", all(got.get(k) == v for k, v in expected.items()), {k: got.get(k) for k in expected})
    p2 = openpyxl.load_workbook(sys.argv[1], data_only=True)["Pareto"]
    gotp = [r[1] for r in p2.iter_rows(min_row=5, max_row=4 + len(order), max_col=2, values_only=True)]
    check("Pareto sau khi tính lại khớp số đếm độc lập", gotp == [cnt[o] for o in order], gotp)
    errs = [c.coordinate for ws in openpyxl.load_workbook(sys.argv[1], data_only=True) for row in ws.iter_rows() for c in row if isinstance(c.value, str) and re.match(r"#(REF|DIV|VALUE|NAME|N/A)", c.value)]
    check("Không có ô lỗi công thức (#REF!, #DIV/0!...)", not errs, errs[:5])

print("== E. Nghiệm thu 2 - action có owner và deadline")
acts = table(wb["Prevention_Actions"], 4)
bad = [V(a, "Action ID") for a in acts for k in ("Action (việc cụ thể)", "Owner Team (đề xuất)", "Owner xác nhận?", "Deadline", "Priority", "Verification (cách xác nhận hoàn thành)", "Status") if V(a, k) in (None, "")]
check("Mọi action có đủ owner, deadline, priority, verification, status (%d action)" % len(acts), not bad, bad)
bad = [V(a, "Action ID") for a in acts if V(a, "Status") != "Done" and (V(a, "Deadline") != "Chưa xác nhận" or V(a, "Owner xác nhận?") != "Chưa xác nhận")]
check("Action chưa được giao ghi rõ 'Chưa xác nhận', không tự đặt tên người/ngày", not bad, bad)
vague = [V(a, "Action ID") for a in acts if re.search(r"cẩn thận|kiểm tra kỹ hơn|chú ý hơn", str(V(a, "Action (việc cụ thể)")).lower()) or len(str(V(a, "Action (việc cụ thể)"))) < 50]
check("Action cụ thể, không chung chung", not vague, vague)
aid = {V(a, "Action ID") for a in acts}
bad = [V(r, "Defect ID") for r in log if V(r, "Defect Kind") != "External" and V(r, "Prevention Action") not in aid]
check("Mọi defect (trừ hệ thống ngoài) gắn với một action tồn tại", not bad, bad)
rca = table(wb["RCA_Top5"], 4)
check("RCA có đúng 5 nhóm, đủ symptom/impact/root cause/detection/corrective/preventive/owner/deadline",
      len(rca) == 5 and all(V(x, k) for x in rca for k in ("Symptom (tester quan sát)", "Impact", "Root cause", "Contributing factors", "Detection", "Corrective action (sửa lỗi hiện tại)", "Preventive action", "Owner", "Deadline")))

print("\nTỔNG: %d/%d PASS" % (sum(res), len(res)))
sys.exit(0 if all(res) else 1)
