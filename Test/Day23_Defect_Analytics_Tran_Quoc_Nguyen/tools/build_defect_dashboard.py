#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Sinh Day23_Defect_Dashboard_Tran_Quoc_Nguyen.xlsx từ tools/defect_data.py.

Mọi con số trên sheet Dashboard và Pareto là CÔNG THỨC COUNTIF/COUNTIFS đọc từ sheet Defect_Log,
nên truy ngược được về từng dòng defect. Chạy: python tools/build_defect_dashboard.py  (cần openpyxl)
"""
import collections, datetime, os, sys
from urllib.parse import quote
import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__)); DAY = os.path.dirname(HERE); TEST = os.path.dirname(DAY)
sys.path.insert(0, HERE)
import defect_data as D

OUT = os.environ.get("DAY23_XLSX") or os.path.join(DAY, "Day23_Defect_Dashboard_Tran_Quoc_Nguyen.xlsx")
ICT = datetime.timezone(datetime.timedelta(hours=7))
NOW = datetime.datetime.now(ICT).strftime("%Y-%m-%d %H:%M")
RETESTED = {"DEF-002", "DEF-003", "DEF-006", "DEF-007", "DEF-012", "DEF-024", "DEF-033", "DEF-036", "DEF-038"}
FIXED_SAME_DAY = {d[0] for d in D.DEFECTS if d[1] == 22 and d[10] == "Fixed"} | {"DEF-038"} | set(D.FIXED_DATES)
VERIFIED = {i: "2026-10-05" for i in RETESTED}; VERIFIED.update(D.LAST_VERIFIED)

HEAD = PatternFill("solid", fgColor="1F3864"); THIN = Side(style="thin", color="BFBFBF"); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINK = Font(color="0563C1", underline="single")
COL = {"Open": "FFC7CE", "In Progress": "FFEB9C", "Fixed": "C6EFCE", "Retest": "FFEB9C", "Closed": "C6EFCE", "Needs verification": "D9D9D9",
       "Critical": "FF7C80", "High": "FFC7CE", "Medium": "FFEB9C", "Low": "DDEBF7", "Done": "C6EFCE", "Proposed": "DDEBF7"}


def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c); x.font = Font(bold=True, color="FFFFFF"); x.fill = HEAD; x.border = BOX
        x.alignment = Alignment(wrap_text=True, vertical="center")


def put(ws, row, vals, color_cols=()):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=row, column=i, value=v); x.border = BOX; x.alignment = Alignment(wrap_text=True, vertical="top")
    for c in color_cols:
        x = ws.cell(row=row, column=c)
        if x.value in COL:
            x.fill = PatternFill("solid", fgColor=COL[x.value])


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=14)
    if sub:
        ws["A2"] = sub; ws["A2"].alignment = Alignment(wrap_text=True)


wb = openpyxl.Workbook()

# ------------------------------------------------------------------ Defect_Log
wl = wb.active; wl.title = "Defect_Log"
LOG_COLS = ["Defect ID", "Source Day", "Component", "Summary", "Symptom (hiện tượng)", "Category", "Severity", "Severity Source", "Priority",
            "Status", "Root Cause (nguyên nhân gốc)", "Root Cause Category", "Detection Phase", "Detection Method", "Owner Team (đề xuất)",
            "Detected Date", "Fixed Date", "Last Verified", "Deadline", "Evidence", "Mở", "Defect Kind", "Fixture Bug", "Regression",
            "Prevention Action", "Ghi chú / kết quả kiểm lại"]
header(wl, 1, LOG_COLS)
for r, d in enumerate(D.DEFECTS, 2):
    (i, day, date, phase, comp, summ, sym, cat, sev, sevsrc, status, root, rc, method, owner, ev, kind, fx, reg, act, note) = d
    rcname = "%s - %s" % (rc, D.RC[rc])
    fixed = "2026-10-05" if i in FIXED_SAME_DAY else ("Not recorded" if status in ("Fixed", "Closed", "Retest") else "")
    put(wl, r, [i, day, comp, summ, sym, cat, sev, sevsrc, D.PRIORITY[sev], status, root, rcname, phase, method, owner, date, fixed,
                VERIFIED.get(i, ""), "Chưa xác nhận" if status not in ("Fixed", "Closed") else "", ev, "Mở", kind, fx, reg, act, note],
        color_cols=(7, 10))
    link = wl.cell(row=r, column=21)
    link.hyperlink = ev if ev.startswith("http") else quote("../" + ev, safe="/.:_-()"); link.font = LINK
N = len(D.DEFECTS); LAST = N + 1
widths(wl, [10, 7, 20, 34, 46, 16, 10, 10, 8, 16, 52, 40, 13, 20, 13, 12, 12, 12, 13, 46, 6, 14, 8, 9, 10, 50])
wl.freeze_panes = "E2"; wl.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(LOG_COLS)), LAST)
L = lambda col: "Defect_Log!$%s$2:$%s$%d" % (col, col, LAST)   # noqa: E731

# ------------------------------------------------------------------ Taxonomy
wt = wb.create_sheet("Taxonomy"); title(wt, "Defect taxonomy dùng thống nhất cho toàn dự án")
r = 3
for name, items, cols in (("Category", D.CATEGORIES, ["Category", "Ví dụ"]), ("Severity", [(s, d + " | Priority mặc định " + D.PRIORITY[s]) for s, d in D.SEVERITIES], ["Severity", "Định nghĩa"]),
                          ("Status", D.STATUSES, ["Status", "Ý nghĩa"]), ("Root Cause Category", [("%s - %s" % kv, "") for kv in D.RC.items()], ["Root Cause Category", ""]),
                          ("Defect Kind", D.KINDS, ["Defect Kind", "Ý nghĩa"]), ("Detection Phase", [(p, "") for p in D.PHASES], ["Detection Phase", ""]),
                          ("Owner Team", [("FE", ""), ("BE", ""), ("FE + BE", ""), ("Data-AI", ""), ("QA", ""), ("Ngoài nhóm", "Hệ thống demo, không thuộc phạm vi sửa")], ["Owner Team", ""])):
    wt.cell(row=r, column=1, value=name).font = Font(bold=True, size=12); r += 1
    header(wt, r, cols); r += 1
    for a, b in items:
        put(wt, r, [a, b]); r += 1
    r += 1
wt.cell(row=r, column=1, value="Quy ước: Symptom = điều tester quan sát được. Root cause = lý do trong code/quy trình khiến hiện tượng xảy ra; không được lặp lại symptom. "
                               "Severity Source: Recorded = lấy từ báo cáo gốc; Proposed = QA đề xuất ngày 05/10 vì báo cáo gốc không ghi.")
widths(wt, [48, 80])

# ------------------------------------------------------------------ Dashboard
wd = wb.create_sheet("Dashboard")
title(wd, "DEFECT DASHBOARD - CyberSoft Learning Hub (Ngày 6 đến Ngày 23)",
      "Sinh lúc %s ICT | Người làm: Trần Quốc Nguyên | Mọi số bên dưới là công thức đếm trên sheet Defect_Log, không nhập tay." % NOW)
wd["A4"] = "KPI"; wd["A4"].font = Font(bold=True, size=12)
kpis = [("Tổng số defect", "=COUNTA(%s)" % L("A")),
        ("Đang mở (Open + In Progress)", '=COUNTIF(%s,"Open")+COUNTIF(%s,"In Progress")' % (L("J"), L("J"))),
        ("Đã sửa/đã đóng (Fixed + Closed)", '=COUNTIF(%s,"Fixed")+COUNTIF(%s,"Closed")' % (L("J"), L("J"))),
        ("Chờ kiểm lại (Retest)", '=COUNTIF(%s,"Retest")' % L("J")),
        ("Cần xác minh (Needs verification)", '=COUNTIF(%s,"Needs verification")' % L("J")),
        ("Critical/High còn mở", '=COUNTIFS(%s,"Open",%s,"Critical")+COUNTIFS(%s,"Open",%s,"High")' % (L("J"), L("G"), L("J"), L("G"))),
        ("Fixture bug", '=COUNTIF(%s,"Y")' % L("W")),
        ("Regression", '=COUNTIF(%s,"Y")' % L("X")),
        ("Lỗi sản phẩm (Defect Kind = Product)", '=COUNTIF(%s,"Product")' % L("V")),
        ("Lỗi do test hoặc môi trường (fixture + tool + environment + CI)", '=COUNTIF(%s,"Test fixture")+COUNTIF(%s,"Test tool")+COUNTIF(%s,"Environment")+COUNTIF(%s,"CI/Process")' % (L("V"), L("V"), L("V"), L("V"))),
        ("Defect phát hiện ở Ngày 23 (test web local 05/10, 07/10 và PR thật)", "=COUNTIF(%s,23)" % L("B")),
        ("Thời gian xử lý trung bình", "Not recorded")]
for i, (k, f) in enumerate(kpis, 5):
    put(wd, i, [k, f]); wd.cell(row=i, column=1).font = Font(bold=True)
wd["A18"] = "Ghi chú: báo cáo gốc không ghi ngày sửa cho đa số lỗi, nên không tính thời gian xử lý. Owner là nhóm đề xuất; chưa có deadline nào được giao chính thức."
wd["A18"].alignment = Alignment(wrap_text=True)


def block(ws, top, col, name, keys, logcol, chart=None, anchor=None, w=14, h=7.5):
    ws.cell(row=top, column=col, value=name).font = Font(bold=True, size=12)
    header_cells = [name.split(" theo ")[-1].capitalize(), "Số defect"]
    for j, t in enumerate(header_cells):
        x = ws.cell(row=top + 1, column=col + j, value=t); x.font = Font(bold=True, color="FFFFFF"); x.fill = HEAD; x.border = BOX
    for i, k in enumerate(keys):
        a = ws.cell(row=top + 2 + i, column=col, value=("Ngày %d" % k) if isinstance(k, int) else k); a.border = BOX
        if k in COL:
            a.fill = PatternFill("solid", fgColor=COL[k])
        crit = '"%s"' % k if not isinstance(k, int) else str(k)
        b = ws.cell(row=top + 2 + i, column=col + 1, value="=COUNTIF(%s,%s)" % (L(logcol), crit)); b.border = BOX
    end = top + 1 + len(keys)
    if chart:
        ch = chart(); ch.title = name; ch.height = h; ch.width = w
        ch.add_data(Reference(ws, min_col=col + 1, min_row=top + 1, max_row=end), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=col, min_row=top + 2, max_row=end))
        if isinstance(ch, BarChart):
            ch.legend = None
        ws.add_chart(ch, anchor)
    return end


used = collections.Counter(d[7] for d in D.DEFECTS)
cats = [c for c, _ in sorted(used.items(), key=lambda kv: -kv[1])] + [c for c, _ in D.CATEGORIES if c not in used]
block(wd, 20, 1, "Defect theo status", [s for s, _ in D.STATUSES], "J", PieChart, "D4", 11, 7.5)
block(wd, 29, 1, "Defect theo severity", [s for s, _ in D.SEVERITIES], "G", BarChart, "K4", 11, 7.5)
block(wd, 36, 1, "Defect theo defect kind", [k for k, _ in D.KINDS], "V", BarChart, "D20", 11, 7.5)
block(wd, 46, 1, "Defect theo category", cats, "F", BarChart, "K20", 16, 9)
block(wd, 67, 1, "Defect theo owner team (đề xuất)", ["FE", "BE", "FE + BE", "Data-AI", "QA", "Ngoài nhóm"], "O", BarChart, "D36", 11, 7.5)
block(wd, 76, 1, "Defect theo source day", sorted({d[1] for d in D.DEFECTS}), "B")
wd["A91"] = "Top 5 root cause và Pareto: xem sheet Pareto và RCA_Top5. Danh sách chi tiết: sheet Defect_Log (lọc theo cột Status/Severity)."
wd["D52"] = "Defect đang mở mức Critical/High"; wd["D52"].font = Font(bold=True, size=12)
for j, t in enumerate(["Defect ID", "Severity", "Summary", "Owner Team", "Action"]):
    x = wd.cell(row=53, column=4 + j, value=t); x.font = Font(bold=True, color="FFFFFF"); x.fill = HEAD; x.border = BOX
rr = 54
for d in D.DEFECTS:
    if d[10] == "Open" and d[8] in ("Critical", "High"):
        row = D.DEFECTS.index(d) + 2
        for j, colL in enumerate(["A", "G", "D", "O", "Y"]):
            x = wd.cell(row=rr, column=4 + j, value='=IF(Defect_Log!%s%d="","",Defect_Log!%s%d)' % (colL, row, colL, row)); x.border = BOX; x.alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1
widths(wd, [44, 12, 3, 12, 10, 52, 12, 10, 3, 3, 12, 12, 12, 12, 12])

# ------------------------------------------------------------------ Pareto
wp = wb.create_sheet("Pareto")
title(wp, "Pareto theo Root Cause Category", "Số lỗi = COUNTIF trên cột Root Cause Category của Defect_Log. Thứ tự sắp xếp được chốt lúc sinh file theo số lỗi giảm dần.")
header(wp, 4, ["Root Cause Category", "Số lỗi", "Tỷ lệ", "Tích lũy", "Còn mở", "Trong top 5 RCA?"])
order = [rc for rc, _ in sorted(collections.Counter(d[12] for d in D.DEFECTS).items(), key=lambda kv: (-kv[1], kv[0]))]
top5 = {x["rc"] for x in D.RCA_TOP5}
for i, rc in enumerate(order):
    r = 5 + i; name = "%s - %s" % (rc, D.RC[rc])
    put(wp, r, [name, '=COUNTIF(%s,A%d)' % (L("L"), r), "=B%d/$B$%d" % (r, 5 + len(order)), "=SUM($B$5:B%d)/$B$%d" % (r, 5 + len(order)),
                '=COUNTIFS(%s,A%d,%s,"Open")' % (L("L"), r, L("J")), "Có" if rc in top5 else ""])
    wp.cell(row=r, column=3).number_format = "0.0%"; wp.cell(row=r, column=4).number_format = "0.0%"
tot = 5 + len(order)
put(wp, tot, ["Tổng", "=SUM(B5:B%d)" % (tot - 1), "=SUM(C5:C%d)" % (tot - 1), "", "=SUM(E5:E%d)" % (tot - 1), ""])
wp.cell(row=tot, column=1).font = Font(bold=True); wp.cell(row=tot, column=3).number_format = "0.0%"
bar = BarChart(); bar.type = "col"; bar.title = "Pareto nguyên nhân gốc"; bar.y_axis.title = "Số lỗi"; bar.height = 10; bar.width = 26
bar.add_data(Reference(wp, min_col=2, min_row=4, max_row=tot - 1), titles_from_data=True)
bar.set_categories(Reference(wp, min_col=1, min_row=5, max_row=tot - 1))
line = LineChart(); line.add_data(Reference(wp, min_col=4, min_row=4, max_row=tot - 1), titles_from_data=True)
line.y_axis.axId = 200; line.y_axis.title = "Tích lũy"; line.y_axis.number_format = "0%"; line.y_axis.scaling.max = 1; line.y_axis.scaling.min = 0; line.y_axis.crosses = "max"
bar += line
wp.add_chart(bar, "A%d" % (tot + 5))
wp.merge_cells(start_row=tot + 2, start_column=1, end_row=tot + 2, end_column=6)
wp.cell(row=tot + 2, column=1, value=D.RCA_SELECTION_NOTE).alignment = Alignment(wrap_text=True, vertical="top")
widths(wp, [62, 9, 9, 10, 9, 16]); wp.row_dimensions[tot + 2].height = 75

# ------------------------------------------------------------------ RCA_Top5
wr = wb.create_sheet("RCA_Top5"); title(wr, "RCA cho 5 nhóm nguyên nhân gốc", D.RCA_SELECTION_NOTE); wr.row_dimensions[2].height = 48
header(wr, 4, ["#", "Root Cause Category", "Nhóm lỗi", "Defect liên quan", "Số lỗi", "Symptom (tester quan sát)", "Impact", "Root cause", "Contributing factors",
               "Detection", "Corrective action (sửa lỗi hiện tại)", "Preventive action", "Owner", "Deadline"])
owners = {a[0]: a[3] for a in D.ACTIONS}
for i, x in enumerate(D.RCA_TOP5, 1):
    r = 4 + i; name = "%s - %s" % (x["rc"], D.RC[x["rc"]])
    pa = x["preventive"].split(",")[-1].strip().split(" ")[0]
    put(wr, r, [i, name, x["group"], x["defects"], '=COUNTIF(%s,B%d)' % (L("L"), r), x["symptom"], x["impact"], x["root"], x["factors"], x["detection"],
                x["corrective"], x["preventive"], owners[pa] + " (đề xuất, chưa xác nhận)", "Chưa xác nhận"])
wr.cell(row=11, column=1, value="Owner và Deadline chưa được quản lý giao chính thức; đây là thông tin cần mentor/trưởng nhóm bổ sung.")
widths(wr, [4, 34, 30, 26, 7, 44, 40, 48, 44, 30, 40, 18, 18, 14])

# ------------------------------------------------------------------ Prevention_Actions
wa = wb.create_sheet("Prevention_Actions"); title(wa, "Prevention actions", "Owner là nhóm đề xuất. Deadline 'Chưa xác nhận' nghĩa là chưa ai giao hạn thật; cần quản lý bổ sung.")
header(wa, 4, ["Action ID", "Root Cause Category", "Action (việc cụ thể)", "Owner Team (đề xuất)", "Owner xác nhận?", "Deadline", "Priority", "Verification (cách xác nhận hoàn thành)", "Status",
               "Số defect liên quan", "Ghi chú"])
for i, a in enumerate(D.ACTIONS, 5):
    aid, rc, act, owner, prio, ver, st, note = a
    put(wa, i, [aid, "%s - %s" % (rc, D.RC[rc]), act, owner, "Chưa xác nhận" if st != "Done" else "QA (đã làm)", "Chưa xác nhận" if st != "Done" else "2026-10-05 (đã xong)", prio, ver, st,
                '=COUNTIF(%s,A%d)' % (L("Y"), i), note], color_cols=(9,))
widths(wa, [10, 36, 66, 14, 15, 18, 8, 56, 11, 9, 40])

# ------------------------------------------------------------------ AI_WORKLOG
ww = wb.create_sheet("AI_WORKLOG"); title(ww, "AI_WORKLOG - Ngày 23")
rows = [
 ("Bài toán trước AI", "Defect nằm rải rác trong báo cáo, junit, workbook và AI_WORKLOG của Ngày 6-22; chưa có taxonomy chung, chưa biết nguyên nhân nào lặp lại nhiều nhất."),
 ("Công cụ đã dùng", "Claude (Cowork): đọc artefact, gọi API web local qua trình duyệt tích hợp, viết generator. Python + openpyxl. Git. GitHub Actions API. Node test runner, pytest (chạy lại ở Ngày 22)."),
 ("Chỉ dẫn chính cho AI", "(1) Bản mô tả yêu cầu Ngày 23: chỉ dùng defect có bằng chứng, phân biệt symptom và root cause, không tự đặt ngày/owner. (2) Chọn 'làm luôn', 'kiểm lỗi cũ và tìm lỗi mới trên web local', 'owner/deadline ghi Chưa xác nhận'."),
 ("AI đề xuất gì", "Taxonomy 18 category, 13 root cause category, 7 defect kind; 58 defect; Pareto; RCA top 5; 14 prevention action; dashboard dùng công thức COUNTIF thay vì nhập số."),
 ("AI sai hoặc chưa đủ ở đâu", "Xem bảng bên dưới."),
 ("Tester kiểm chứng thế nào", "Chạy tools/verify_defect_dashboard.py (đếm lại từ Defect_Log, kiểm evidence tồn tại, kiểm symptom khác root cause); tự mở 3 defect bất kỳ và mở file evidence. (Bạn điền kết quả chạy trên máy mình.)"),
 ("Quyết định của tester", "(Bạn tự điền: đồng ý/không đồng ý top 5; severity nào cần đổi; lỗi nào loại khỏi log.)"),
]
header(ww, 3, ["Mục", "Nội dung"])
for i, (a, b) in enumerate(rows, 4):
    put(ww, i, [a, b]); ww.cell(row=i, column=1).font = Font(bold=True)
r = 4 + len(rows) + 1
ww.cell(row=r, column=1, value="AI sai/chưa đủ và cách xử lý").font = Font(bold=True, size=12); r += 1
header(ww, r, ["#", "AI làm/đề xuất", "Sai hoặc chưa đủ ở đâu", "Kiểm bằng gì", "Xử lý"]); r += 1
mistakes = [
 ("Định đưa 3 lỗi của bản dashboard Ngày 21 (CI ghi '0 evidence', ngưỡng '= 100%' thành công thức, data quality sai phạm vi) vào log", "Thư mục Day21 có AI_WORKLOG ghi các lỗi đó đã bị thay bằng bản khác, không còn file bằng chứng", "Kiểm đường dẫn evidence: không tồn tại", "Loại cả 3 khỏi Defect_Log; chỉ giữ DEF-038 vì có bằng chứng từ GitHub Actions"),
 ("Xếp test_tampered_token_is_401 (Ngày 7, 41/42) là lỗi phân quyền của sản phẩm", "Bộ test chạy trên API mô phỏng; chạy lại ngày 05/10 cho 42/42", "pytest trong thư mục Day07", "Đổi thành Test Automation, Status = Needs verification, không tính vào lỗi sản phẩm"),
 ("Coi F-D13-06 (judge trim khoảng trắng cuối) là defect", "Báo cáo gốc ghi mức Info và là chính sách so sánh, không có expected result bị vi phạm", "Đọc sheet 06_Findings", "Không đưa vào log"),
 ("Gộp 13 case FAIL của Ngày 19 thành 13 defect", "10 case RT-041..050 cùng một hiện tượng và một nguyên nhân; đếm 10 lần sẽ làm lệch Pareto", "Đọc cột Actual Response: 10 dòng giống nhau", "Ghi 1 defect (DEF-033) kèm số case; 3 case còn lại là 3 defect riêng"),
 ("Ghi DEF-054 (forgot-password 500) với root cause 'thiếu cấu hình SMTP'", "Chưa xem log BE nên mới là giả thuyết", "Chỉ có response 500", "Status = Needs verification; root cause ghi rõ là nghi ngờ"),
 ("Ghi chéo sai mã trong ghi chú (DEF-052 thay vì DEF-053)", "Đánh số lại sau khi thêm lỗi", "Script kiểm số lượng và đọc lại", "Đã sửa ghi chú"),
 ("Test trực tiếp tạo dữ liệu trong DB local (1 tài khoản email sai định dạng, 1 lượt quiz, 1 bài nộp)", "Không dọn được qua API", "Ghi trong evidence mục C", "Báo cho tester để dọn; không tính là defect"),
]
for i, m in enumerate(mistakes, 1):
    put(ww, r, [i] + list(m)); r += 1
widths(ww, [30, 70, 60, 40, 60])

# ------------------------------------------------------------------ README
wre = wb.create_sheet("README", 0); title(wre, "Ngày 23 - Defect analytics và root cause")
kind_n = collections.Counter(d[16] for d in D.DEFECTS)
info = [
 ("Mục đích", "Biết loại lỗi nào lặp lại, nguyên nhân gốc là gì và cần đổi quy trình/automated test thế nào để lỗi không quay lại."),
 ("Phạm vi dữ liệu", "Defect có bằng chứng từ Ngày 7 đến Ngày 22 của kế hoạch QA, cộng lỗi phát hiện khi test web local ngày 05/10 và 07/10/2026 (Ngày 23). Ngày 1-6 và Ngày 16-17 không có defect có bằng chứng."),
 ("Nguồn dữ liệu", "JUnit Ngày 7; báo cáo và regression Ngày 8; workbook Ngày 9, 10, 12, 13, 14, 19, 20; lint-report Ngày 11; AI_WORKLOG Ngày 11, 12, 15, 22; báo cáo Ngày 18; GitHub Actions; evidence/live_test_2026-10-05.md và live_test_2026-10-07.md."),
 ("Cách dùng", "Dashboard: xem tổng quan. Pareto + RCA_Top5: nguyên nhân gốc. Prevention_Actions: việc cần làm. Defect_Log: bấm cột 'Mở' để mở file bằng chứng (đường dẫn tương đối, cần giữ file trong thư mục Test/Day23...)."),
 ("Số liệu", "Mọi con số là công thức đếm trên Defect_Log. Tại thời điểm sinh file: %d defect, trong đó %d lỗi sản phẩm." % (len(D.DEFECTS), kind_n["Product"])),
 ("Không đưa vào log", "Lỗi cố ý gài để thực hành (seeded/demo) của Ngày 6, 7, 14; finding trên file mẫu cố ý sai của Ngày 11, 12; các mục CANDIDATE chưa kiểm của Ngày 14; ca test PASS; 3 lỗi của bản dashboard Ngày 21 vì không còn file bằng chứng."),
 ("Giới hạn dữ liệu", "Ngày phát hiện của Ngày 9, 10 không được ghi trong artefact (Not recorded). Hầu hết lỗi không có ngày sửa nên không tính thời gian xử lý. Severity của nhiều lỗi do QA đề xuất ngày 05/10 (cột Severity Source = Proposed). Owner là nhóm đề xuất, chưa ai được giao; Deadline = Chưa xác nhận."),
 ("Giới hạn test web", "DB local khác bộ seed trong repo. Ngày 05/10 chỉ test qua API và DOM. Ngày 07/10 xem được giao diện ở chiều rộng 769 px nhưng trang FE trong trình duyệt tích hợp không gọi được BE, nên không test luồng cần đăng nhập qua giao diện. Chi tiết trong 2 file evidence."),
 ("Sinh lại file", "python tools/build_defect_dashboard.py rồi python tools/verify_defect_dashboard.py. Sửa dữ liệu trong tools/defect_data.py, không sửa tay trong Excel."),
 ("Thời điểm sinh", NOW + " ICT"),
]
header(wre, 3, ["Mục", "Nội dung"])
for i, (a, b) in enumerate(info, 4):
    put(wre, i, [a, b]); wre.cell(row=i, column=1).font = Font(bold=True)
widths(wre, [24, 130])

wb._sheets = [wb[n] for n in ["README", "Defect_Log", "Taxonomy", "Dashboard", "Pareto", "RCA_Top5", "Prevention_Actions", "AI_WORKLOG"]]
wb.save(OUT)
print("Wrote", os.path.basename(OUT), "|", len(D.DEFECTS), "defects |", len(D.ACTIONS), "actions | pareto order:", order)
